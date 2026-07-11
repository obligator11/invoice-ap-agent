from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass

from PIL import Image

from ingestion.pdf_parser import extract_native_pdf_text, render_pdf_pages_as_images
from ingestion.ocr import ocr_image, ocr_pdf_pages
from schemas.models import SourceMethod

SUPPORTED_OFFICE_EXTENSIONS = {".docx", ".xlsx"}


@dataclass
class IngestResult:
    raw_text: str
    source_method: SourceMethod
    low_confidence: bool = False


def ingest_file(filepath: str, original_filename: str | None = None) -> IngestResult:
    name = (original_filename or filepath).lower()

    if name.endswith(".pdf"):
        return _ingest_pdf(filepath)
    if name.endswith((".jpg", ".jpeg", ".png")):
        return _ingest_image(filepath)
    if name.endswith(".docx"):
        return _ingest_docx(filepath)
    if name.endswith(".xlsx"):
        return _ingest_xlsx(filepath)

    raise ValueError(f"Unsupported file type: {name}")


def ingest_bytes(content: bytes, filename: str) -> IngestResult:
    """Used for email attachments, which arrive as bytes rather than a
    path already on disk."""
    suffix = os.path.splitext(filename)[1] or ".bin"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name
    try:
        result = ingest_file(tmp_path, original_filename=filename)
    finally:
        os.unlink(tmp_path)

    # An email attachment that turns out to be a native PDF is still
    # tagged as email_attachment, not native_pdf — we want to know the
    # *channel* it arrived through, independent of its internal format.
    result.source_method = SourceMethod.EMAIL_ATTACHMENT
    return result


def _ingest_pdf(filepath: str) -> IngestResult:
    result = extract_native_pdf_text(filepath)
    if not result.needs_ocr:
        return IngestResult(raw_text=result.text, source_method=SourceMethod.NATIVE_PDF)

    images = render_pdf_pages_as_images(filepath)
    ocr_result = ocr_pdf_pages(images)
    return IngestResult(
        raw_text=ocr_result.text,
        source_method=SourceMethod.OCR_SCAN,
        low_confidence=ocr_result.low_confidence,
    )


def _ingest_image(filepath: str) -> IngestResult:
    image = Image.open(filepath)
    ocr_result = ocr_image(image)
    return IngestResult(
        raw_text=ocr_result.text,
        source_method=SourceMethod.OCR_PHOTO,
        low_confidence=ocr_result.low_confidence,
    )


def _ingest_docx(filepath: str) -> IngestResult:
    import docx  # python-docx
    doc = docx.Document(filepath)
    text = "\n".join(p.text for p in doc.paragraphs)
    return IngestResult(raw_text=text, source_method=SourceMethod.OFFICE_DOC)


def _ingest_xlsx(filepath: str) -> IngestResult:
    import openpyxl
    wb = openpyxl.load_workbook(filepath, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lines.append(" ".join(str(c) for c in row if c is not None))
    return IngestResult(raw_text="\n".join(lines), source_method=SourceMethod.OFFICE_DOC)
